from __future__ import annotations

import csv
import io
from pathlib import Path

import fitz
import pandas as pd
import pdfplumber
from docx import Document

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff"}
SPREADSHEET_EXTENSIONS = {".xlsx", ".xls", ".csv"}
TEXT_EXTENSIONS = {".tex", ".txt"}
SUPPORTED_EXTENSIONS = {".pdf", ".docx", *IMAGE_EXTENSIONS, *SPREADSHEET_EXTENSIONS, *TEXT_EXTENSIONS}


def clean_cell(value: object) -> str:
    return "" if value is None else str(value)


def valid_tables(tables: list[list[list[object]]] | None) -> list[list[list[str]]]:
    result = []
    for table in tables or []:
        normalized = [[clean_cell(cell) for cell in row] for row in table]
        if len(normalized) >= 2 and max((len(row) for row in normalized), default=0) >= 2:
            result.append(normalized)
    return result


def ocr_image_bytes(data: bytes, engine) -> str:
    result = engine(data)
    if isinstance(result, tuple):
        result = result[0]
    if not result:
        return ""
    ordered = sorted(result, key=lambda item: (min(point[1] for point in item[0]), min(point[0] for point in item[0])))
    return "\n".join(str(item[1]) for item in ordered if len(item) > 1 and item[1]).strip()


def extract_pdf(data: bytes, engine) -> list[dict]:
    pages = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        if len(pdf.pages) > 50:
            raise ValueError("PDF exceeds the 50 page limit")
        fallback = fitz.open(stream=data, filetype="pdf")
        try:
            for number, page in enumerate(pdf.pages, start=1):
                text = (page.extract_text() or "").strip()
                tables = valid_tables(page.extract_tables())
                if not tables:
                    tables = valid_tables(page.extract_tables({"vertical_strategy": "text", "horizontal_strategy": "text"}))
                if not tables:
                    try:
                        tables = valid_tables([table.extract() for table in fallback[number - 1].find_tables().tables])
                    except Exception:
                        tables = []
                method = "native"
                if len(text) < 20:
                    text = ocr_image_bytes(fallback[number - 1].get_pixmap(dpi=200, alpha=False).tobytes("png"), engine)
                    method = "ocr"
                pages.append({"page": number, "text": text, "tables": tables, "method": method})
        finally:
            fallback.close()
    return pages


def extract_docx(data: bytes) -> list[dict]:
    document = Document(io.BytesIO(data))
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    tables = [[[clean_cell(cell.text) for cell in row.cells] for row in table.rows] for table in document.tables]
    return [{"page": 1, "text": "\n".join(paragraphs), "tables": valid_tables(tables), "method": "native"}]


def extract_spreadsheet(data: bytes, extension: str) -> list[dict]:
    if extension == ".csv":
        rows = list(csv.reader(io.StringIO(data.decode("utf-8", errors="ignore"))))
        return [{"page": 1, "text": "\n".join("\t".join(row) for row in rows), "tables": [rows] if rows else [], "method": "native"}]
    sheets = pd.read_excel(io.BytesIO(data), sheet_name=None, header=None)
    pages = []
    for number, (_, frame) in enumerate(sheets.items(), start=1):
        rows = [[clean_cell(value) for value in row] for row in frame.fillna("").values.tolist()]
        pages.append({"page": number, "text": "\n".join("\t".join(row) for row in rows), "tables": [rows] if rows else [], "method": "native"})
    return pages


def extract_document(filename: str, data: bytes, engine) -> list[dict]:
    extension = Path(filename).suffix.lower()
    if extension == ".pdf": return extract_pdf(data, engine)
    if extension in IMAGE_EXTENSIONS: return [{"page": 1, "text": ocr_image_bytes(data, engine), "tables": [], "method": "ocr"}]
    if extension == ".docx": return extract_docx(data)
    if extension in SPREADSHEET_EXTENSIONS: return extract_spreadsheet(data, extension)
    if extension in TEXT_EXTENSIONS: return [{"page": 1, "text": data.decode("utf-8", errors="ignore"), "tables": [], "method": "native"}]
    raise LookupError("Unsupported file type")
